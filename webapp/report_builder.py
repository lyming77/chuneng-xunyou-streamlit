from __future__ import annotations

import io
import math
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "报告模板/投资分析报告_标准格式.docx"
FONT_BODY = "仿宋_GB2312"
FONT_HEAD = "黑体"
FONT_SECOND = "楷体_GB2312"
FONT_TITLE = "STSong"


def _number(value: float) -> str:
    value = float(value)
    return str(int(value)) if value.is_integer() else f"{value:g}"


def _fmt(value, decimals=2, suffix="") -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    if not math.isfinite(number):
        return "—"
    return f"{number:,.{decimals}f}{suffix}"


def _clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_font(run, name=FONT_BODY, size=16, bold=False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_text(document, text="", *, font=FONT_BODY, size=16, bold=False,
              align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, keep=False):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(32 if indent else 0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(28.95)
    paragraph.paragraph_format.keep_with_next = keep
    _set_font(paragraph.add_run(text), font, size, bold)
    return paragraph


def _add_heading(document, text: str, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(32)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(28.95)
    paragraph.paragraph_format.keep_with_next = True
    _set_font(paragraph.add_run(text), FONT_HEAD if level == 1 else FONT_SECOND, 16, level == 2)
    return paragraph


def _set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(18)
    _set_font(paragraph.add_run(str(text)), FONT_HEAD if bold else FONT_BODY, 10.5, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", 60), ("start", 70), ("bottom", 60), ("end", 70)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def _add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, header in enumerate(headers):
        _set_cell(table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        cant_split = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(cant_split)
        for index, value in enumerate(row):
            _set_cell(cells[index], value,
                      align=WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return table


def _new_document() -> Document:
    document = Document(TEMPLATE) if TEMPLATE.exists() else Document()
    _clear_body(document)
    normal = document.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(16)
    normal.paragraph_format.first_line_indent = Pt(32)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(28.95)
    return document


def _add_cover(document, config, metrics):
    _add_text(document, "", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    _add_text(document, config["project_name"], font=FONT_TITLE, size=22,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, keep=True)
    _add_text(document, f"{_number(metrics['P'])}MW/{_number(metrics['E'])}MWh配建储能项目",
              font=FONT_TITLE, size=22, align=WD_ALIGN_PARAGRAPH.CENTER,
              indent=False, keep=True)
    _add_text(document, "投资分析报告", font=FONT_TITLE, size=22,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, keep=True)
    _add_text(document, "（方案A测算）", font=FONT_SECOND, size=16,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    today = date.today()
    _add_text(document, f"{today.year}年{today.month}月{today.day}日", size=16,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_report_bytes(config: dict, grid: pd.DataFrame, detail: dict) -> bytes:
    metrics = detail["metrics"]
    finance = detail["finance"]
    investment = finance["capex_breakdown"]
    document = _new_document()
    _add_cover(document, config, metrics)

    _add_heading(document, "一、报告摘要")
    _add_table(document, ["核心指标", "测算结果", "单位"], [
        ["初始总投资", _fmt(metrics["capex"], 1), "万元"],
        ["首年经营净收益", _fmt(metrics["R_year"], 1), "万元"],
        ["首年容量补偿", _fmt(metrics["R_cap"], 1), "万元"],
        ["首年总收入", _fmt(metrics["total_revenue"], 1), "万元"],
        ["全投资NPV", _fmt(metrics["fi_npv"], 1), "万元"],
        ["全投资IRR", _fmt(metrics["fi_irr"], 2), "%"],
        ["资本金NPV", _fmt(metrics["eq_npv"], 1), "万元"],
        ["资本金IRR", _fmt(metrics["eq_irr"], 2), "%"],
    ], widths=[6.2, 5.0, 3.2])
    _add_text(document,
              f"本报告依据{metrics['start']:%Y年%m月%d日}至{metrics['end']:%Y年%m月%d日}的全年"
              f"15分钟生产与市场数据，对{_number(metrics['P'])}MW/{_number(metrics['E'])}MWh"
              f"储能方案进行方案A历史回测。测算结果仅陈述当前输入数据、调度规则和经济参数下的结果。")

    _add_heading(document, "二、项目与输入条件")
    _add_table(document, ["项目", "输入值"], [
        ["新能源类型", config["project_type"]],
        ["新能源装机", f"{config['station_power_mw']:g}MW"],
        ["厂用电功率", f"{config['aux_power_mw']:g}MW"],
        ["时序点数", f"{metrics['points']:,}点"],
        ["储能效率", f"{config['efficiency_pct']:g}%"],
        ["SOC运行区间", f"{config['soc_min_pct']:g}%—{config['soc_max_pct']:g}%"],
        ["最低预测价差", f"{config['minimum_spread_yuan_per_mwh']:g}元/MWh"],
        ["最高预测充电价格", f"{config['maximum_charge_price_yuan_per_mwh']:g}元/MWh"],
        ["项目评价期", f"{config['project_years']}年"],
    ], widths=[6.2, 8.2])
    _add_text(document,
              "调度计划使用日前预测电价和日前预测发电功率形成；运行日按实际发电能力修正，"
              "结算采用实时市场价格和实际执行电量。标准建设边界为站内储能及接入工程，不含"
              "新建送出线路和主变扩容。")

    _add_heading(document, "三、储能运行结果")
    _add_table(document, ["指标", "测算值"], [
        ["全年充电量", f"{metrics['charge_mwh']:,.1f}MWh"],
        ["限电电量充电", f"{metrics['curtailed_charge_mwh']:,.1f}MWh"],
        ["偏差弥补电量", f"{metrics['supplemental_charge_mwh']:,.1f}MWh"],
        ["付费发电充电", f"{metrics['paid_charge_mwh']:,.1f}MWh"],
        ["全年放电量", f"{metrics['discharge_mwh']:,.1f}MWh"],
        ["等效满循环", f"{metrics['annual_cycles']:,.1f}次/年"],
        ["限电电量利用率", f"{metrics['curtailment_capture_pct']:,.2f}%"],
        ["付费充电平均成本", f"{_fmt(metrics['average_paid_charge_price'], 2)}元/MWh"],
        ["平均放电收益", f"{_fmt(metrics['average_discharge_value'], 2)}元/MWh"],
    ], widths=[7.0, 7.4])
    _add_text(document,
              f"全年总充电量为{metrics['charge_mwh']:,.1f}MWh，其中限电电量充电"
              f"{metrics['curtailed_charge_mwh']:,.1f}MWh、偏差弥补"
              f"{metrics['supplemental_charge_mwh']:,.1f}MWh、付费发电充电"
              f"{metrics['paid_charge_mwh']:,.1f}MWh。按全年充电量除以储能容量计算，"
              f"等效满循环为{metrics['annual_cycles']:,.1f}次。")

    _add_heading(document, "四、配储前后运行指标")
    before_curtailment_rate = (metrics["curtailed_energy_mwh"] /
                               (metrics["actual_generation_mwh"] + metrics["curtailed_energy_mwh"]) * 100
                               if metrics["actual_generation_mwh"] + metrics["curtailed_energy_mwh"] else 0)
    after_curtailment = max(0.0, metrics["curtailed_energy_mwh"] - metrics["curtailed_charge_mwh"])
    after_curtailment_rate = (after_curtailment /
                              (metrics["actual_generation_mwh"] + metrics["curtailed_energy_mwh"]) * 100
                              if metrics["actual_generation_mwh"] + metrics["curtailed_energy_mwh"] else 0)
    _add_table(document, ["指标", "建设前", "建设后", "变化"], [
        ["整体上网电量(MWh)", _fmt(metrics["actual_generation_mwh"], 1),
         _fmt(metrics["post_export_mwh"], 1),
         _fmt(metrics["post_export_mwh"] - metrics["actual_generation_mwh"], 1)],
        ["限电量(MWh)", _fmt(metrics["curtailed_energy_mwh"], 1), _fmt(after_curtailment, 1),
         _fmt(-metrics["curtailed_charge_mwh"], 1)],
        ["限电比例(%)", _fmt(before_curtailment_rate, 2), _fmt(after_curtailment_rate, 2),
         _fmt(after_curtailment_rate - before_curtailment_rate, 2)],
        ["实时市场平均结算电价(元/MWh)", _fmt(metrics["market_price_before"], 2),
         _fmt(metrics["market_price_after"], 2),
         _fmt(metrics["market_price_after"] - metrics["market_price_before"], 2)],
    ], widths=[5.0, 3.2, 3.2, 3.2])

    _add_heading(document, "五、投资构成")
    labels = [
        ("电池系统", "battery_system"), ("土地及场地", "land"), ("PCS", "pcs"),
        ("站内升压、开关、保护及消防", "station_equipment"),
        ("建筑安装工程", "construction_install"), ("前期研究及申报", "predevelopment"),
        ("建设管理、勘察设计及监理", "management_design_supervision"),
        ("基本预备费", "contingency"),
    ]
    _add_table(document, ["投资项目", "金额(万元)", "占比"], [
        [label, _fmt(investment[key] / 1e4, 1),
         _fmt(investment[key] / finance["capex"] * 100 if finance["capex"] else 0, 1, "%")]
        for label, key in labels
    ] + [["合计", _fmt(finance["capex"] / 1e4, 1), "100.0%"]], widths=[7.4, 3.6, 3.4])
    _add_text(document,
              f"初始总投资为{finance['capex']/1e4:,.1f}万元。造价包括站内储能设备、场地、"
              "建筑安装、前期工作、设计监理和基本预备费，不包括新建送出线路和主变扩容。")

    _add_heading(document, "六、收益与财务评价")
    _add_table(document, ["收益项目", "万元/年"], [
        ["抵厂用电收益", _fmt(metrics["rev_tou_wan"], 1)],
        ["市场放电收入", _fmt(metrics["rev_market_wan"], 1)],
        ["减：充电成本", _fmt(metrics["charging_cost_wan"], 1)],
        ["经营净收益", _fmt(metrics["R_year"], 1)],
        ["容量补偿", _fmt(metrics["R_cap"], 1)],
        ["首年总收入", _fmt(metrics["total_revenue"], 1)],
    ], widths=[8.0, 6.4])
    _add_text(document,
              f"全投资NPV为{metrics['fi_npv']:,.1f}万元，全投资IRR为{_fmt(metrics['fi_irr'], 2)}%；"
              f"资本金NPV为{metrics['eq_npv']:,.1f}万元，资本金IRR为{_fmt(metrics['eq_irr'], 2)}%。"
              f"评价期末税后净回收价值为{metrics['terminal_net_wan']:,.1f}万元。")

    _add_heading(document, "七、全网格位置")
    successful = grid[grid["error"].fillna("").eq("")].copy() if "error" in grid else grid.copy()
    same_duration = successful[successful["H"].eq(metrics["H"])].sort_values("P")
    display = same_duration[["P", "E", "annual_cycles", "total_revenue", "capex", "fi_npv", "eq_irr"]]
    if len(display) > 12:
        distances = (display["P"] - metrics["P"]).abs()
        display = display.loc[distances.nsmallest(12).index].sort_values("P")
    _add_table(document,
               ["功率(MW)", "容量(MWh)", "循环次数", "首年收入(万元)", "投资(万元)",
                "全投资NPV(万元)", "资本金IRR(%)"],
               [[_fmt(row.P, 1), _fmt(row.E, 1), _fmt(row.annual_cycles, 1),
                 _fmt(row.total_revenue, 1), _fmt(row.capex, 1), _fmt(row.fi_npv, 1),
                 _fmt(row.eq_irr, 2)] for row in display.itertuples()],
               widths=[2.1, 2.2, 2.2, 2.5, 2.4, 2.8, 2.5])
    _add_text(document,
              f"上表列示{metrics['H']}小时配置在全网格中的相邻容量结果，用于比较容量扩大后"
              "循环次数、收入、投资和资本金收益率的变化。")

    _add_heading(document, "八、主要测算口径")
    items = [
        "方案A使用日前预测电价和日前预测发电功率形成96点计划，实际执行按实时发电能力修正。",
        f"预测充电价格不高于{config['maximum_charge_price_yuan_per_mwh']:g}元/MWh，"
        f"预测充放价差不低于{config['minimum_spread_yuan_per_mwh']:g}元/MWh。",
        f"SOC运行区间为{config['soc_min_pct']:g}%—{config['soc_max_pct']:g}%，"
        f"储能综合效率为{config['efficiency_pct']:g}%。",
        f"项目评价期{config['project_years']}年，资本金比例{config['equity_ratio_pct']:g}%，"
        f"贷款利率{config['loan_rate_pct']:g}%，贷款期限{config['loan_term_years']}年。",
        "限电量只用于限电指标计算，不作为充电资格限制。",
    ]
    for index, item in enumerate(items, 1):
        _add_text(document, f"{index}. {item}")

    properties = document.core_properties
    properties.title = f"{config['project_name']}储能投资分析报告"
    properties.subject = "配储寻优方案A测算"
    properties.author = ""
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
