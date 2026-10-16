# -*- coding: utf-8 -*-
"""读取标准收资表格,校验并生成寻优标准输入。

用法: python3 脚本/读取收资表格.py [表格路径]
默认读根目录 收资表格.xlsx。
输出: 时序 DataFrame + 参数字典 + λ_tou(分时电价),并写 CSV。
"""
import sys
import pandas as pd
from openpyxl import load_workbook

DEFAULT = "收资表格.xlsx"
COEF_DEFAULT = {"尖峰": 2.0, "高峰": 1.7, "平段": 1.0, "低谷": 0.3, "深谷": 0.1}
MONTH_GROUP = {1: "1-2月、12月", 2: "1-2月、12月", 12: "1-2月、12月",
               3: "3-5月", 4: "3-5月", 5: "3-5月",
               6: "6月", 7: "7-8月", 8: "7-8月",
               9: "9-11月", 10: "9-11月", 11: "9-11月"}


def read_kv(path, sheet):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    return {str(r[0]).strip(): r[1] for r in rows[1:] if r[0] not in (None, "")}


def read_ts(path):
    df = pd.read_excel(path, sheet_name="时序数据")
    df.columns = [c.strip() for c in df.columns]
    return df


def check_ts(df):
    assert len(df) == 365 * 96, f"时序数据行数 {len(df)},应为 {365*96}"
    t = pd.to_datetime(df["时间戳"])
    assert (t.diff().dt.total_seconds().iloc[1:] == 900).all(), "时间戳非 15min 连续"
    req = ["实际发电功率(MW)", "超短期预测功率(MW)", "预测发电功率(日前,MW)",
           "实时市场价格(元/MWh)", "预测价格(元/MWh)"]
    for c in req:
        assert df[c].notna().all(), f"必填列「{c}」存在空值"
    print(f"时序校验通过: {len(df)} 行,15min 连续,必填列完整")
    return t


def parse_ranges(s):
    if s in (None, "", "—", "-"):
        return []
    out = []
    for seg in str(s).split(","):
        seg = seg.strip()
        if "-" in seg:
            a, b = seg.split("-")
            out.append([float(a), float(b)])
    return out


def read_tou_table(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["时段表"].iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() for h in rows[0]]
    zones = {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        zones[str(r[0]).strip()] = {header[j]: parse_ranges(r[j]) for j in range(1, len(header))}
    return zones


def zone_of(zones, month, k):
    h = k / 4.0
    for cat, ranges in zones[MONTH_GROUP[month]].items():
        for a, b in ranges:
            if a <= h < b:
                return cat
    return "平段"


def build_tou(flat, coef, zones, t):
    df = pd.DataFrame({"time": t})
    df["month"] = df.time.dt.month
    df["k"] = df.groupby(df.time.dt.normalize()).cumcount()
    df["tou"] = [flat[m] * coef[zone_of(zones, m, k)] for m, k in zip(df["month"], df["k"])]
    return df[["time", "tou"]]


def main(path=DEFAULT):
    params = {
        "场站": read_kv(path, "场站参数"),
        "储能": read_kv(path, "储能参数"),
        "技术经济": read_kv(path, "技术经济参数"),
    }
    flat_raw = read_kv(path, "平段电价")
    coef_raw = read_kv(path, "峰谷平系数")
    flat = {int(str(k).replace("月", "")): v for k, v in flat_raw.items()}
    coef = {k: float(v) for k, v in coef_raw.items()} if coef_raw else COEF_DEFAULT

    assert params["场站"]["装机容量"] not in (None, ""), "装机容量必填"
    assert params["场站"]["场站类型"] in ("光伏", "风电", "混合"), "场站类型必填(光伏/风电/混合)"
    assert params["场站"]["厂用电功率"] not in (None, ""), "厂用电功率必填"
    assert params["场站"]["机制电量/上网电量比 R"] in (0, 80), "机制电量/上网电量比 R 需为 0(主动)或 80(被动)"
    assert params["技术经济"]["容量补偿电价 P容量"] not in (None, ""), "容量补偿电价 P容量 必填"
    fin_keys = ("电池造价 c_电池", "土地造价 c_土地", "PCS 造价 c_PCS", "其他造价 c_其他",
                "资本金比例", "贷款利率", "贷款期限", "资本金回报率",
                "电池折旧年限", "交流侧折旧年限", "电池年衰减率", "残值率",
                "维护费用", "所得税率", "项目年限 n")
    missing = [k for k in fin_keys if params["技术经济"].get(k) in (None, "")]
    assert not missing, f"技术经济参数缺必填项: {missing}"
    assert len(flat) == 12 and all(v is not None for v in flat.values()), "平段电价需 12 个月齐备"

    df = read_ts(path)
    t = check_ts(df)
    zones = read_tou_table(path)
    tou = build_tou(flat, coef, zones, t)

    print("场站:", params["场站"])
    print("储能:", params["储能"])
    print("技术经济:", params["技术经济"])
    print("平段电价(元/MWh):", {k: v for k, v in sorted(flat.items())})
    print("λ_tou 范围: %.1f ~ %.1f 元/MWh" % (tou["tou"].min(), tou["tou"].max()))
    return df, params, tou


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    df, params, tou = main(path)
    out = df.rename(columns={
        "时间戳": "time", "实际发电功率(MW)": "p_gen", "超短期预测功率(MW)": "p_avail",
        "理论功率(MW)": "p_theo", "预测发电功率(日前,MW)": "p_fcst",
        "实时市场价格(元/MWh)": "mkt_price", "预测价格(元/MWh)": "fcst_price"})
    out.to_csv("15min数据/收资_时序.csv", index=False, encoding="utf-8-sig")
    tou.to_csv("15min数据/收资_分时电价.csv", index=False, encoding="utf-8-sig")
    print("已输出 15min数据/收资_时序.csv 与 15min数据/收资_分时电价.csv")
